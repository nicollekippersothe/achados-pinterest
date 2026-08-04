"""
Extrai texto de slides (.pptx) e preenche templates HTML automaticamente.
Uso: python slide_extractor.py arquivo.pptx
"""

import sys
import json
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    print("Instale: pip install python-pptx")
    sys.exit(1)

try:
    from jinja2 import Template
except ImportError:
    print("Instale: pip install jinja2")
    sys.exit(1)


# Template HTML para pin do Pinterest
PIN_TEMPLATE = """<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{{ titulo }} — Achados BR</title>
<style>
  * { margin: 0; padding: 0; box-sizing: border-box; }
  body { font-family: Georgia, serif; background: #F2EDE4; display: flex; justify-content: center; align-items: center; min-height: 100vh; }
  .pin { width: 400px; background: #fff; border-radius: 16px; overflow: hidden; box-shadow: 0 4px 24px rgba(0,0,0,0.10); }
  .pin-header { background: #7A8C6E; padding: 32px 24px 24px; text-align: center; }
  .pin-header h1 { font-size: 13px; letter-spacing: 4px; color: #fff; font-weight: 400; }
  .pin-body { padding: 28px 24px; }
  .produto { font-size: 22px; color: #3D3530; line-height: 1.4; margin-bottom: 16px; }
  .descricao { font-size: 14px; color: #5a5248; line-height: 1.7; margin-bottom: 20px; }
  .preco { font-size: 28px; color: #7A8C6E; font-weight: 700; margin-bottom: 8px; }
  .preco-de { font-size: 14px; color: #9C8E80; text-decoration: line-through; margin-bottom: 20px; }
  .btn { display: block; background: #E63946; color: #fff; text-align: center; padding: 14px; border-radius: 8px; text-decoration: none; font-size: 14px; letter-spacing: 2px; }
  .tags { margin-top: 16px; display: flex; flex-wrap: wrap; gap: 6px; }
  .tag { font-size: 11px; color: #7A8C6E; background: #F2EDE4; padding: 4px 10px; border-radius: 20px; }
  .pin-footer { padding: 12px 24px; border-top: 1px solid #F2EDE4; }
  .marca { font-size: 11px; letter-spacing: 3px; color: #9C8E80; }
</style>
</head>
<body>
<div class="pin">
  <div class="pin-header"><h1>ACHADOS BR</h1></div>
  <div class="pin-body">
    <div class="produto">{{ titulo }}</div>
    {% if descricao %}
    <div class="descricao">{{ descricao }}</div>
    {% endif %}
    {% if preco_de %}
    <div class="preco-de">De {{ preco_de }}</div>
    {% endif %}
    {% if preco %}
    <div class="preco">{{ preco }}</div>
    {% endif %}
    {% if link %}
    <a class="btn" href="{{ link }}" target="_blank">VER OFERTA →</a>
    {% endif %}
    {% if tags %}
    <div class="tags">
      {% for tag in tags %}<span class="tag">#{{ tag }}</span>{% endfor %}
    </div>
    {% endif %}
  </div>
  <div class="pin-footer"><div class="marca">ACHADOS BR · BRASIL</div></div>
</div>
</body>
</html>"""


def extrair_preco(texto):
    """Detecta padrões de preço no texto: R$ 49,90 / R$49.90 / 49,90"""
    padrao = r'R?\$?\s*(\d{1,4}[.,]\d{2})'
    match = re.search(padrao, texto, re.IGNORECASE)
    if match:
        valor = match.group(0).strip()
        if not valor.startswith('R'):
            valor = 'R$ ' + valor
        return valor
    return None


def extrair_link(texto):
    """Detecta URLs no texto."""
    padrao = r'https?://[^\s]+'
    match = re.search(padrao, texto)
    return match.group(0) if match else None


def extrair_slide(caminho_pptx):
    """Extrai dados estruturados de cada slide do arquivo."""
    prs = Presentation(caminho_pptx)
    slides_data = []

    for i, slide in enumerate(prs.slides, 1):
        textos = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    texto = para.text.strip()
                    if texto:
                        textos.append(texto)

        if not textos:
            continue

        # Heurística: primeiro texto = título, demais = descrição
        titulo = textos[0]
        descricao_partes = []
        preco = None
        preco_de = None
        link = None
        tags = []

        for texto in textos[1:]:
            p = extrair_preco(texto)
            l = extrair_link(texto)
            if p and preco is None:
                preco = p
            elif p and preco is not None and preco_de is None:
                preco_de = preco
                preco = p
            elif l:
                link = l
            elif texto.startswith('#'):
                tags += [t.lstrip('#') for t in texto.split() if t.startswith('#')]
            else:
                descricao_partes.append(texto)

        slides_data.append({
            'slide': i,
            'titulo': titulo,
            'descricao': ' '.join(descricao_partes) or None,
            'preco': preco,
            'preco_de': preco_de,
            'link': link,
            'tags': tags,
        })

    return slides_data


def gerar_pins(caminho_pptx, pasta_saida='pins'):
    """Extrai slides e gera um HTML por slide."""
    Path(pasta_saida).mkdir(exist_ok=True)
    template = Template(PIN_TEMPLATE)

    dados = extrair_slide(caminho_pptx)
    if not dados:
        print("Nenhum texto encontrado nos slides.")
        return

    json_path = Path(pasta_saida) / 'dados.json'
    json_path.write_text(json.dumps(dados, ensure_ascii=False, indent=2))
    print(f"Dados extraídos salvos em: {json_path}")

    for slide in dados:
        html = template.render(**slide)
        nome = f"pin_slide_{slide['slide']:02d}.html"
        saida = Path(pasta_saida) / nome
        saida.write_text(html, encoding='utf-8')
        print(f"Pin gerado: {saida}  |  Título: {slide['titulo'][:50]}")

    print(f"\n{len(dados)} pin(s) gerado(s) em '{pasta_saida}/'")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Uso: python slide_extractor.py arquivo.pptx [pasta_saida]")
        print("\nExemplo de teste com slide simulado:")

        # Demonstração sem arquivo real
        dados_exemplo = {
            'slide': 1,
            'titulo': 'Kit Skincare Completo',
            'descricao': 'Hidratante + sérum + protetor solar. Ideal para pele seca.',
            'preco': 'R$ 89,90',
            'preco_de': 'R$ 149,90',
            'link': 'https://s.shopee.com.br/exemplo',
            'tags': ['skincare', 'beleza', 'oferta'],
        }

        Path('pins').mkdir(exist_ok=True)
        template = Template(PIN_TEMPLATE)
        html = template.render(**dados_exemplo)
        saida = Path('pins') / 'pin_exemplo.html'
        saida.write_text(html, encoding='utf-8')
        print(f"Pin de exemplo gerado: {saida}")
        sys.exit(0)

    caminho = sys.argv[1]
    pasta = sys.argv[2] if len(sys.argv) > 2 else 'pins'
    gerar_pins(caminho, pasta)
